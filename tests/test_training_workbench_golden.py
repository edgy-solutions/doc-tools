"""Gate for the training-workbench contract (docs/contracts/training-workbench.md).

One test per surface, all against ``tests/fixtures/training/golden.json``. A
surface that a later lane-tr step delivers is ``xfail(strict=True)`` naming that
step; the step's PR flips it (strict, so an unflipped pass fails the suite).
The ``nondeterministic`` block of the golden is LLM-derived example data: only
its SHAPE is asserted, by ``assert_llm_shape``, which later live or fake-driven
tests reuse.
"""
import hashlib
import json
import re
from pathlib import Path
from types import SimpleNamespace

import pytest

from doc_tools.utils.concept_identity import concept_node_id, normalize_concept_name
from doc_tools.utils.content_kind import KIND_MAPPING

REPO_ROOT = Path(__file__).resolve().parents[1]
FIXTURES = Path(__file__).parent / "fixtures" / "training"
DECK = FIXTURES / "hydraulics_fixture.pptx"
META = FIXTURES / "metadata.json"
GOLDEN = json.loads((FIXTURES / "golden.json").read_text(encoding="utf-8"))

INGEST_ID = GOLDEN["ingest_id"]
HEX = INGEST_ID.removeprefix("sha256:")
SLIDES = GOLDEN["neo4j"]["slides"]["nodes"]
SLIDE_IDS = [s["id"] for s in SLIDES]
NONDET = GOLDEN["neo4j"]["nondeterministic"]


def assert_llm_shape(sections, concepts, teaches, slide_ids, ingest_id):
    """Shape of LLM-derived graph data (sections, concepts, TEACHES rows)."""
    numbers = {int(i.rsplit("_p", 1)[1]) for i in slide_ids}
    names = [c["name"] for c in concepts]
    # A Concept's identity is (domain, normalized name) and `id` is its stable key
    # (contract section 3). So the uniqueness that matters is on `id`, not on `name`:
    # asserting unique names would pass on a fixture whose two nodes normalize alike,
    # which is exactly the duplicate the id constraint exists to refuse. And the id is
    # asserted RECOMPUTABLE from the node's own stored fields, so a fixture (or a node
    # in a live store) cannot carry an id that nothing can re-derive.
    ids = [c["id"] for c in concepts]
    assert len(ids) == len(set(ids)), "one Concept per (domain, normalized name)"
    for c in concepts:
        assert c["domain"] == "TRAINING", c
        assert c["name_normalized"] == normalize_concept_name(c["name"]), c
        assert c["id"] == concept_node_id(c["domain"], c["name"]), c
    sec_ids = {s["id"] for s in sections}
    id_re = re.compile(rf"^{re.escape(ingest_id)}(_s\d+)+$")
    for s in sections:
        assert id_re.match(s["id"]), s["id"]
        assert s["parent"] == ingest_id or s["parent"] in sec_ids, s["parent"]
        assert isinstance(s["level"], int) and s["level"] >= 1
        assert s["start_page"] in numbers, s["start_page"]
        assert isinstance(s["end_page"], int) and s["end_page"] >= 0
        assert len(s["concept_summary"]) <= 10
        assert s.get("summary") is None or (isinstance(s["summary"], str) and s["summary"]), s["id"]
        assert isinstance(s["objectives"], list) and all(isinstance(o, str) for o in s["objectives"])
        assert all(n in names for n in s["concept_summary"]), s["concept_summary"]
    for t in teaches:
        assert t["slide"] in slide_ids, t
        assert t["concept"] in names, t
        assert isinstance(t["salience"], float) and 0.0 <= t["salience"] <= 1.0, t
        assert isinstance(t["related_terms"], list) and all(isinstance(r, str) for r in t["related_terms"]), t


def test_fixture_bytes_are_pinned():
    fx = GOLDEN["fixture"]
    assert hashlib.sha256(DECK.read_bytes()).hexdigest() == fx["deck_sha256"]
    assert hashlib.sha256(META.read_bytes()).hexdigest() == fx["metadata_sha256"]
    assert GOLDEN["ingest_id"] == "sha256:" + fx["deck_sha256"]


def test_golden_ids_follow_the_contract():
    from weaviate.util import generate_uuid5

    course = GOLDEN["neo4j"]["course"]["props"]
    assert course["id"] == INGEST_ID
    assert course["current_key"] == f"training/generated/{HEX}/current.json"
    assert course["title"] == json.loads(META.read_text(encoding="utf-8"))["title"]
    assert course["filename"] == "hydraulics_fixture.pptx"

    numbers = [s["number"] for s in SLIDES]
    assert all(isinstance(n, int) for n in numbers)
    assert numbers == list(range(1, len(numbers) + 1))
    for s in SLIDES:
        assert s["id"] == f"{INGEST_ID}_p{s['number']}"
        assert s["id"].rsplit("_p", 1) == [INGEST_ID, str(s["number"])]
        assert len(s["text"]) <= 500
        assert not any(re.match(r"^\[\w+\] $", ln) for ln in s["text"].split("\n")), s["id"]

    mi = GOLDEN["minio"]
    prefix = mi["prefix"]
    assert prefix == f"training/generated/{HEX}/<pipeline_version>/"
    cj = mi["current_json"]
    assert cj["prefix"] == prefix and prefix.endswith("/")
    assert cj["manifest_key"] == prefix + "manifest.json"
    keys = mi["keys_required"]
    assert all(k.startswith(prefix) for k in keys)
    for k in ("text.json", "manifest.json", *(f"pages/page_{n}.png" for n in numbers)):
        assert prefix + k in keys, k

    objs = GOLDEN["weaviate"]["objects"]
    assert len(objs) == len(SLIDES)
    by_chunk = {o["properties"]["chunk_id"]: o for o in objs}
    for s in SLIDES:
        o = by_chunk[s["id"]]
        p = o["properties"]
        assert p["doc_id"] == INGEST_ID and p["domain"] == "TRAINING"
        assert o["uuid"] == str(generate_uuid5(p["chunk_id"]))
        assert p["text"][:500] == s["text"]
    for bad in GOLDEN["weaviate"]["forbidden_chunk_id_substrings"]:
        assert not any(bad in c for c in by_chunk), bad


def test_nondeterministic_example_has_the_contract_shape():
    assert_llm_shape(NONDET["sections"]["nodes"], NONDET["concepts"]["nodes"],
                     NONDET["teaches"], SLIDE_IDS, INGEST_ID)


def test_example_edges_follow_the_section_4_2_rules():
    """The explicit `edges` block is what wb's fake store loads, so it must be
    exactly what the contract's rules give from the example sections and TEACHES."""
    sections, edges = NONDET["sections"]["nodes"], NONDET["edges"]
    assert edges["HAS_SECTION"] == [{"from": s["parent"], "to": s["id"]} for s in sections]
    assert edges["TEACHES"] == [{"from": t["slide"], "to": t["concept"], "salience": t["salience"],
                                 "related_terms": t["related_terms"]}
                                for t in NONDET["teaches"]]
    numbers = {s["id"]: s["number"] for s in SLIDES}
    covers = []
    for sec in sections:
        reached = [sid for sid, n in numbers.items()
                   if n >= sec["start_page"] and (sec["end_page"] == 0 or n <= sec["end_page"])]
        assert sorted(e["to"] for e in edges["section_slide"] if e["from"] == sec["id"]) == sorted(reached)
        sal = {}
        for t in NONDET["teaches"]:
            if t["slide"] in reached:
                sal.setdefault(t["concept"], []).append(t["salience"])
        mine = sorted(({"from": sec["id"], "to": c, "score": round(sum(v) / len(v), 6), "frequency": len(v)}
                       for c, v in sal.items()), key=lambda r: -r["score"])
        assert sec["concept_summary"] == [r["to"] for r in mine][:10], sec["id"]
        covers += mine
    assert sorted(edges["COVERS"], key=json.dumps) == sorted(covers, key=json.dumps)
    # ALIAS_OF (contract 4.2): both endpoints are Concepts of the example, the status
    # is a candidate, and the edge is exactly what related_terms gives under the rule.
    names = {c["name"]: c for c in NONDET["concepts"]["nodes"]}
    want_alias = []
    for t in NONDET["teaches"]:
        for term in t["related_terms"]:
            pair = {"from": t["concept"], "to": term}
            if (term in names and normalize_concept_name(term) != normalize_concept_name(t["concept"])
                    and not any(a["from"] == pair["from"] and a["to"] == pair["to"] for a in want_alias)):
                want_alias.append({**pair, "status": "candidate", "source": "related_terms",
                                   "course_id": INGEST_ID})
    assert edges["ALIAS_OF"] == want_alias and want_alias
    assert all(a["from"] in names and a["to"] in names for a in edges["ALIAS_OF"])
    assert "ALIAS_OF" not in GOLDEN["neo4j"]["never_written"]["relations"]

    # Section objects: one per Section with notes, keyed by the Section id.
    from weaviate.util import generate_uuid5
    sec_ids = {s["id"]: s for s in sections}
    objs = NONDET["section_objects"]["objects"]
    assert objs
    assert {o["properties"]["chunk_id"] for o in objs} == \
        {s["id"] for s in sections if s.get("summary") or s["objectives"]}
    for o in objs:
        p = o["properties"]
        sec = sec_ids[p["chunk_id"]]
        assert o["uuid"] == str(generate_uuid5(p["chunk_id"]))
        assert p["chunk_kind"] == "section" and p["section_id"] == p["chunk_id"]
        assert p["page"] == sec["start_page"] and p["heading"] == sec["title"]
        assert p["doc_id"] == INGEST_ID and p["domain"] == "TRAINING"
    assert all(o["properties"]["chunk_kind"] == "slide" for o in GOLDEN["weaviate"]["objects"])


def _strip(el, keys):
    el = json.loads(json.dumps(el, default=str))
    for k in keys:
        head, _, tail = k.partition(".")
        if tail:
            el.get(head, {}).pop(tail, None)
        else:
            el.pop(head, None)
    return el


def test_extraction_reproduces_the_golden_slides(tmp_path):
    from collections import defaultdict

    from doc_tools.utils.extraction import extract_text_and_metadata
    from doc_tools.utils.layout_detector import detect_layout
    from doc_tools.utils.pptx_media_extractor import extract_images_from_pptx

    elements = extract_text_and_metadata(str(DECK), extract_images=True,
                                         image_output_dir=str(tmp_path))
    by_page = defaultdict(list)
    for el in elements:
        by_page[el.get("metadata", {}).get("page_number", 1)].append(el)
    vol = GOLDEN["volatile_element_keys"]
    got = []
    for n in sorted(by_page):
        els = by_page[n]
        text = "\n".join(f"[{e.get('type', 'Text')}] {e.get('text', '')}"
                         for e in els if e.get("text", "").strip())
        got.append({"id": f"{INGEST_ID}_p{n}", "number": n, "text": text[:500],
                    "asset_type": "PPTX", "layout_style": detect_layout(els),
                    "elements": [_strip(e, vol) for e in els]})
    assert got == SLIDES

    media = sorted(Path(p).name for p in extract_images_from_pptx(str(DECK), str(tmp_path)))
    prefix = GOLDEN["minio"]["prefix"] + "images/"
    want = sorted(k[len(prefix):] for k in GOLDEN["minio"]["keys_required"] if k.startswith(prefix))
    assert media == want


def test_never_written_lists_the_workbench_owned_names():
    nw = GOLDEN["neo4j"]["never_written"]
    assert set(nw["labels"]) == {"Project", "TargetNode", "User"}
    assert set(nw["relations"]) == {"OWNS", "SUGGESTED_SOURCE", "DERIVED_FROM"}
    assert GOLDEN["neo4j"]["written_only_by"] == {
        "concepts.normalise": {"labels": ["CanonicalConcept"],
                               "relations": ["ALIGNS_TO", "PROPOSED_ALIGNS_TO"]}}
    banned = [":Project", ":TargetNode", ":SUGGESTED_SOURCE", ":DERIVED_FROM", ":OWNS"]
    # Contract v5 section 4.4/4.7: only `concepts.normalise` (doc_tools/passes/concepts.py)
    # may name these; everywhere else they stay banned.
    only_concepts = [":CanonicalConcept", ":ALIGNS_TO", ":PROPOSED_ALIGNS_TO"]
    sanctioned = Path("doc_tools") / "passes" / "concepts.py"
    hits = []
    for py in (REPO_ROOT / "doc_tools").rglob("*.py"):
        src = py.read_text(encoding="utf-8", errors="replace")
        rel = py.relative_to(REPO_ROOT)
        names = banned + ([] if rel == sanctioned else only_concepts)
        hits += [f"{rel}: {b}" for b in names if b in src]
    assert not hits, hits


def test_training_kind_row_and_sensor():
    row = KIND_MAPPING.get("course")
    assert row is not None, "no KIND_MAPPING row 'course'"
    assert row.kind == "course" and row.domain_type == "training"
    assert {".pptx", ".pdf", ".docx"} <= {s if s.startswith(".") else "." + s
                                          for s in row.file_suffixes}
    # ingress_prefix is the Lane-1 ingress-user/<fmt>/ path (sealed to {pdf, xml}
    # by test_ingress_prefix_registry.py); the course row must leave it None.
    assert row.ingress_prefix is None
    src = (REPO_ROOT / "doc_tools" / "definitions.py").read_text(encoding="utf-8")
    assert "training/course/" in src
    # Contract section 2: the PATH carries the kind (ADR-0021 <domain>/<kind>/...),
    # and the sidecar carries neither. The parser stamps domain_type from the
    # first key segment; path-derive then reads "course" from the second.
    from doc_tools.utils.content_kind import resolve_content_kind

    meta = json.loads(META.read_text(encoding="utf-8"))
    assert "content_kind" not in meta and "domain_type" not in meta
    deck_key = GOLDEN["drop"]["keys_in_order"][-1]
    assert deck_key.startswith("training/course/")
    assert resolve_content_kind({**meta, "domain_type": "training"}, deck_key) is row


def test_teaches_salience_is_on_the_edge():
    from doc_tools.plugins.models import BaseSection, DocumentNode
    from doc_tools.plugins.training import Concept, SlideAugmentation, TrainingPlugin

    sid = f"{INGEST_ID}_p2"
    rows = [t for t in NONDET["teaches"] if t["slide"] == sid]
    assert rows
    aug = SlideAugmentation(concepts=[Concept(name=t["concept"], salience=t["salience"])
                                      for t in rows])
    node = DocumentNode(
        base_extraction=BaseSection(title="s2", level=1, page_start=2, content="x", node_id=sid),
        domain_augmentation=aug)
    config = SimpleNamespace(graph_node_label="Document", graph_child_label="Slide")
    cypher, _ = TrainingPlugin(domain_type="training").to_graph_queries(
        [node], config, doc_id=INGEST_ID)

    concept_merge = re.compile(r"MERGE\s*\(\s*\w*\s*:Concept[^{)]*\{([^}]*)\}")
    teaches_rel = re.compile(r"\[\s*\w*\s*:TEACHES\b([^\]]*)\]")
    # the MERGEs that write the edge, not the stale-edge prune that follows them
    edge_queries = [q for q in cypher
                    if teaches_rel.search(q["query"]) and "DELETE t" not in q["query"]]
    assert edge_queries, "no query writes TEACHES"
    for q in edge_queries:
        for m in concept_merge.finditer(q["query"]):
            keys = re.findall(r"(\w+)\s*:", m.group(1))
            assert "salience" not in keys, "salience in the Concept MERGE key"
            # `id`, not `name`, since the architect's 2026-10-24 ruling: the id is
            # hash(domain, normalized name), which is one property and so can carry the
            # uniqueness constraint Neo4j had on nothing. The contract no longer says
            # "one Concept:TRAINING per name across courses" -- that wording called
            # `name` the identity, and the identity is (domain, normalized name). The
            # guarantee is one node per NORMALIZED name per domain, which is strictly
            # stronger: two spellings that normalize alike now MERGE onto one node
            # where merging on `name` left two.
            assert keys == ["id"], keys
        assert "TRAINING" in q["query"] or "TRAINING" in map(str, q["params"].values())
        assert re.search(r"\.domain\s*=|domain\s*:", q["query"]), "Concept domain not set"

    def sets_salience_on_edge(q):
        rel = teaches_rel.search(q["query"])
        after = q["query"][rel.start():]
        return bool(re.search(r"\{[^}]*salience\s*:", rel.group(0))
                    or re.search(r"\w+\.salience\s*=", after))

    for t in rows:
        assert any(t["salience"] in q["params"].values() and sets_salience_on_edge(q)
                   for q in edge_queries), t


PIPELINE_VERSION = "doc-tools@steptest"
BUCKET = GOLDEN["drop"]["bucket"]


def _sub(value):
    """Golden placeholders -> this run's pipeline version."""
    return value.replace("<pipeline_version>", PIPELINE_VERSION)


def _parse_drop(client, monkeypatch, key=None, body=None):
    """Run the real parser (real extraction) on the fixture drop in a fake S3."""
    from dagster import build_asset_context
    from dag_tools.components.s3_sensor.file_component import S3FileConfig

    from doc_tools.components.document_parser import DocumentParserComponent
    from tests.test_document_parser_versioning import _FakeS3Resource

    monkeypatch.setenv("DOC_TOOLS_VERSION", PIPELINE_VERSION)
    # Page PNGs are lane tr step e's; the PDF rasteriser has nothing to do on a deck.
    monkeypatch.setattr("doc_tools.components.document_parser.rasterize_pdf_pages",
                        lambda *a, **k: [])
    meta_key, deck_key = GOLDEN["drop"]["keys_in_order"]
    client.put_object(Bucket=BUCKET, Key=meta_key, Body=META.read_bytes())
    client.put_object(Bucket=BUCKET, Key=deck_key, Body=DECK.read_bytes())
    defs = DocumentParserComponent(
        name="process_document_artifact", partition_name="pdf_files").build_defs(None)
    asset_def = next(iter(defs.assets))
    ctx = build_asset_context(partition_key=deck_key)
    cfg = S3FileConfig(file_url=f"s3://{BUCKET}/{deck_key}")
    return asset_def(ctx, config=cfg, s3=_FakeS3Resource(client))


class FakeWeaviate:
    """Weaviate resource whose collection records data.insert/replace kwargs."""

    def __init__(self, existing=()):
        from unittest.mock import MagicMock

        self.inserts, self.replaces, self.existing = [], [], set(existing)
        data = MagicMock()
        data.exists.side_effect = lambda u: u in self.existing
        data.insert.side_effect = lambda **kw: self.inserts.append(kw)
        data.replace.side_effect = lambda **kw: self.replaces.append(kw)
        self.data = data
        self.client = MagicMock()
        self.client.collections.get.return_value = MagicMock(data=data)

    def get_client(self):
        return self.client

    @property
    def objects(self):
        return self.inserts + self.replaces


def _graph_run(manifest, client, monkeypatch, augment_calls=None, weaviate=None,
               reraise=False, fulltext=None, neo4j_out=None, extract=None, normalise=None,
               ctx_out=None):
    """Run the graph asset with fake neo4j/weaviate and a stubbed plugin/LLM.

    Returns (neo4j execute_query calls as (query, params), chunk_ids indexed).
    With `weaviate=` (a FakeWeaviate) the REAL _index_chunk runs against it;
    otherwise _index_chunk is stubbed and `indexed` collects its properties.
    The default concept call (`_extract_section`) RAISES, so with `reraise=True` the run fails on
    failed:concept_parse; pass `extract=` (self, section_text, instructions) to reach a later
    failure. `augment_calls` collects the slide numbers each call carried."""
    from unittest.mock import MagicMock

    from dagster import build_asset_context

    from doc_tools.assets import semantic_assets as sa
    from doc_tools.config import IngestionConfig
    from doc_tools.plugins.training import TrainingPlugin
    from tests.test_document_parser_versioning import _FakeS3Resource

    indexed = []
    # The per-course concepts.normalise call is stubbed: it reads the whole domain from the
    # graph, which this harness fakes. `normalise=` (a callable) replaces the stub.
    monkeypatch.setattr("doc_tools.passes.concepts.normalise", normalise or (lambda *a, **k: {}))
    monkeypatch.setattr(sa, "_ensure_weaviate_collection", lambda *a, **k: None)
    if weaviate is None:
        monkeypatch.setattr(sa, "_index_chunk",
                            lambda _c, _n, props, **_k: indexed.append(props) or "written")

    def _extract(self, section_text, instructions):
        if augment_calls is not None:
            augment_calls.extend(int(n) for n in re.findall(r"=== Slide (\d+) ===", section_text))
        if extract is not None:
            return extract(self, section_text, instructions)
        raise RuntimeError("plugin/LLM stubbed out")

    monkeypatch.setattr(TrainingPlugin, "_extract_section", _extract)
    monkeypatch.setattr(TrainingPlugin, "process_fulltext", fulltext or (lambda *a, **k: []))
    neo4j = MagicMock()
    if neo4j_out is not None:
        neo4j_out.append(neo4j)
    config = IngestionConfig(graph_node_label="Document", graph_child_label="Chunk",
                             vector_collection_name="DocumentChunk", bucket=BUCKET)
    try:
        ctx = build_asset_context()
        if ctx_out is not None:
            ctx_out.append(ctx)
        sa.build_knowledge_graph(
            ctx, config, manifest, s3=_FakeS3Resource(client),
            neo4j=neo4j, weaviate=weaviate or MagicMock(), llm=MagicMock(), jena=MagicMock())
    except Exception:  # noqa: BLE001 -- the writes under test precede any later stage
        if reraise:
            raise
    calls = neo4j.get_client.return_value.execute_query.call_args_list
    return [(c.args[0], c.args[1]) for c in calls], indexed


def _labels(query, var):
    m = re.search(rf"\({var}((?::\w+)+)\s*\{{id", query)
    return set(m.group(1).split(":")[1:])


def test_ids_and_versioned_layout(monkeypatch):
    from tests.test_document_parser_versioning import _FakeS3Client

    client = _FakeS3Client()
    manifest = _parse_drop(client, monkeypatch)
    mi = GOLDEN["minio"]

    # Pages (`pages/page_<n>.png`) are lane tr step e's; everything else is step c's.
    written = {k for (_b, k) in client.write_order}
    for key in mi["keys_required"]:
        if "/pages/" in key:
            continue
        assert _sub(key) in written, key

    current_key = mi["current_json_key"]
    assert current_key in written
    pointer = json.loads(client.objects[(BUCKET, current_key)])
    want = {k: _sub(v) for k, v in mi["current_json"].items()}
    from datetime import datetime
    datetime.fromisoformat(pointer.pop("updated_at"))
    want.pop("updated_at")
    assert pointer == want

    assert manifest["doc_id"] == manifest["ingest_id"] == INGEST_ID
    assert manifest["current_key"] == current_key
    stored = json.loads(client.objects[(BUCKET, _sub(mi["prefix"]) + "manifest.json")])
    assert stored["doc_id"] == stored["ingest_id"] == INGEST_ID
    assert stored["current_key"] == current_key
    # Nothing landed under the old filename-keyed layout.
    assert not any("/generated/hydraulics_fixture_pptx/" in k for k in written)

    # --- the graph asset, fed that manifest and the text.json it wrote ---
    calls, indexed = _graph_run(manifest, client, monkeypatch)
    parent_q, parent_p = calls[0]
    assert _labels(parent_q, "n") | set(re.findall(r"SET n:(\w+)", parent_q)) == \
        set(GOLDEN["neo4j"]["course"]["labels"])
    for k in ("id", "title", "filename", "current_key"):
        assert parent_p[k] == GOLDEN["neo4j"]["course"]["props"][k], k

    slide_calls = [(q, p) for q, p in calls if "MERGE (child:" in q]
    assert [p["id"] for _q, p in slide_calls] == [s["id"] for s in SLIDES]
    for (q, p), golden in zip(slide_calls, SLIDES):
        assert _labels(q, "child") == {"Chunk", "Slide", "TRAINING"}
        assert p["text"] == golden["text"]
        for edge in GOLDEN["neo4j"]["course_slide_edges"]:
            assert f"MERGE (parent)-[:{edge}]->(child)" in q, edge
    assert [o["chunk_id"] for o in indexed] == SLIDE_IDS


def test_a_text_less_slide_still_gets_a_slide_node_and_no_weaviate_object(monkeypatch):
    from tests.test_document_parser_versioning import _FakeS3Client

    cid = "sha256:" + "b" * 64
    text_key = f"training/generated/{'b' * 64}/{PIPELINE_VERSION}/text.json"
    elements = [
        {"type": "Title", "text": "Has text", "metadata": {"page_number": 1}},
        {"type": "Image", "text": "  ", "metadata": {"page_number": 2}},
        {"type": "PageBreak", "text": "", "metadata": {"page_number": 2}},
    ]
    client = _FakeS3Client()
    client.put_object(Bucket=BUCKET, Key=text_key, Body=json.dumps(elements).encode())
    manifest = {"doc_id": cid, "ingest_id": cid, "filename": "empty_slide.pptx",
                "text_location": text_key, "metadata": {"domain_type": "training"},
                "current_key": f"training/generated/{'b' * 64}/current.json"}
    augmented = []
    calls, indexed = _graph_run(manifest, client, monkeypatch, augment_calls=augmented)

    # No metadata.json title: the Course title falls back to the filename stem.
    assert calls[0][1]["title"] == "empty_slide"
    slides = {p["id"]: p for q, p in calls if "MERGE (child:" in q}
    assert list(slides) == [f"{cid}_p1", f"{cid}_p2"]
    assert slides[f"{cid}_p1"]["text"] == "[Title] Has text"
    assert slides[f"{cid}_p2"]["text"] == ""
    assert [o["chunk_id"] for o in indexed] == [f"{cid}_p1"]
    assert set(augmented) == {1}  # the call is retried once, so slide 1 is seen twice


def test_a_non_training_drop_keeps_the_filename_keyed_layout(monkeypatch):
    from dagster import build_asset_context
    from dag_tools.components.s3_sensor.file_component import S3FileConfig

    from doc_tools.components.document_parser import DocumentParserComponent
    from tests.test_document_parser_versioning import (
        _FakeS3Client, _FakeS3Resource, _fake_extract_text_and_metadata,
        _fake_rasterize_pdf_pages,
    )

    monkeypatch.setattr("doc_tools.components.document_parser.extract_text_and_metadata",
                        _fake_extract_text_and_metadata)
    monkeypatch.setattr("doc_tools.components.document_parser.rasterize_pdf_pages",
                        _fake_rasterize_pdf_pages)
    monkeypatch.setattr("doc_tools.utils.crop_geometry.repair_table_crops", lambda *a, **k: 0)
    monkeypatch.setenv("DOC_TOOLS_VERSION", PIPELINE_VERSION)
    key = "sustainment/inbound/testnotice/Test_Notice.pdf"
    client = _FakeS3Client()
    asset_def = next(iter(DocumentParserComponent(
        name="process_document_artifact", partition_name="pdf_files").build_defs(None).assets))
    manifest = asset_def(build_asset_context(partition_key=key),
                         config=S3FileConfig(file_url=f"s3://{BUCKET}/{key}"),
                         s3=_FakeS3Resource(client))

    gen = "sustainment/inbound/testnotice/generated/Test_Notice_pdf"
    written = {k for (_b, k) in client.write_order}
    assert {f"{gen}/{PIPELINE_VERSION}/manifest.json", f"{gen}/{PIPELINE_VERSION}/text.json",
            f"{gen}/current.json"} <= written
    assert not any(k.startswith("training/") for k in written)
    assert manifest["doc_id"] == "Test_Notice"
    assert manifest["current_key"] == f"{gen}/current.json"
    pointer = json.loads(client.objects[(BUCKET, f"{gen}/current.json")])
    assert pointer["prefix"] == f"{gen}/{PIPELINE_VERSION}/"


def test_render_pages_writes_page_pngs(monkeypatch, tmp_path):
    """Every pages/page_{n}.png in golden minio.keys_required is rendered and listed.

    LibreOffice and Poppler are stubbed at the subprocess layer; pages() itself
    (rename, dimensions, 1..N check) runs for real."""
    import os

    from pptx import Presentation
    from PIL import Image

    from doc_tools.passes import render
    from tests.test_document_parser_versioning import _FakeS3Client

    n_slides = len(Presentation(str(DECK)).slides)
    assert n_slides == 4

    monkeypatch.setattr(render, "_convert_to_pdf", lambda src, work_dir, t: str(tmp_path / "x.pdf"))

    def _fake_rasterize(pdf_path, work_dir, dpi, timeout_s):
        out = []
        for n in range(1, n_slides + 1):
            p = os.path.join(work_dir, f"p-{n}.png")
            Image.new("RGB", (20, 10), "white").save(p)
            out.append(p)
        return out

    monkeypatch.setattr(render, "_rasterize", _fake_rasterize)

    client = _FakeS3Client()
    content_types = {}
    orig_upload = client.upload_file

    def _recording_upload(Filename, Bucket, Key, ExtraArgs=None):
        content_types[Key] = (ExtraArgs or {}).get("ContentType")
        return orig_upload(Filename=Filename, Bucket=Bucket, Key=Key, ExtraArgs=ExtraArgs)

    client.upload_file = _recording_upload
    manifest = _parse_drop(client, monkeypatch)

    want = {_sub(k) for k in GOLDEN["minio"]["keys_required"] if "/pages/" in k}
    assert want, "golden lists no pages/ keys"
    written_pages = {k for (_b, k) in client.write_order if "/pages/" in k}
    assert written_pages == want
    for k in want:
        assert content_types[k] == "image/png", k
    assert [p["page"] for p in manifest["pages"]] == list(range(1, n_slides + 1))
    assert {p["key"] for p in manifest["pages"]} == want
    assert "pages_error" not in manifest


def test_vectors_written_and_no_figure_objects(monkeypatch):
    from tests.test_document_parser_versioning import _FakeS3Client

    monkeypatch.setattr("doc_tools.utils.embed.embed_document", lambda text: [0.5] * 768)
    client = _FakeS3Client()
    manifest = _parse_drop(client, monkeypatch)
    assert manifest.get("embedded_images"), "fixture drop must carry an embedded image"
    weaviate = FakeWeaviate()
    _graph_run(manifest, client, monkeypatch, weaviate=weaviate)

    ids = [o["properties"]["chunk_id"] for o in weaviate.objects]
    assert ids == SLIDE_IDS
    for bad in GOLDEN["weaviate"]["forbidden_chunk_id_substrings"]:
        assert not any(bad in cid for cid in ids), bad
    assert weaviate.inserts and not weaviate.replaces
    for o in weaviate.inserts:
        assert len(o["vector"]) == 768


def test_outline_pages_come_from_element_metadata(monkeypatch):
    """Section pages are element page_number, never the LLM's: with the pinned
    (wrong-paged) ExtractOutline reply stubbed in, the real pass reproduces the
    golden's sections."""
    from types import SimpleNamespace

    from doc_tools.plugins.training import TrainingPlugin

    elements = [{**el, "metadata": {**el.get("metadata", {}), "page_number": s["number"]}}
                for s in SLIDES for el in s["elements"]]
    fallback = GOLDEN["neo4j"]["outline"]["fallback"]

    def _obj(n):
        return SimpleNamespace(title=n["title"], level=n["level"], start_page=n["start_page"],
                               end_page=n["end_page"],
                               subsections=[_obj(c) for c in n["subsections"]])

    plugin = TrainingPlugin(domain_type="training")
    monkeypatch.setattr(TrainingPlugin, "execute_global_pass",
                        lambda self, els, max_chars: ([SimpleNamespace(
                            sections=[_obj(n) for n in fallback])], []))
    nodes = plugin.process_fulltext("", INGEST_ID, {}, elements=elements)
    aug = nodes[0].domain_augmentation
    assert aug.outline_source == GOLDEN["neo4j"]["outline"]["source"] == "llm"
    assert GOLDEN["neo4j"]["course"]["props"]["outline_source"] == "llm"

    got = []

    def flat(secs, parent):
        for i, s in enumerate(secs):
            sid = f"{parent}_s{i}"
            got.append({"id": sid, "title": s.title, "level": s.level,
                        "start_page": s.start_page, "end_page": s.end_page, "parent": parent})
            flat(s.subsections, sid)

    flat(aug.sections, INGEST_ID)
    want = [{k: s[k] for k in ("id", "title", "level", "start_page", "end_page", "parent")}
            for s in NONDET["sections"]["nodes"]]
    assert got == want

    llm_pages = {}

    def collect(ns):
        for n in ns:
            llm_pages[n["title"]] = (n["start_page"], n["end_page"])
            collect(n["subsections"])

    collect(fallback)
    differing = [s for s in got if llm_pages.get(s["title"]) not in (None, (s["start_page"], s["end_page"]))]
    assert differing, "the stub's pages must differ from the derived ones for this test to bite"


def empty_section_reply(section_text):
    """A concept-call reply with an entry for every slide in the text and no concepts."""
    from doc_tools.baml_client.types import SectionAugmentation, SlideConcepts

    return SectionAugmentation(
        slides=[SlideConcepts(slide=int(n), concepts=[], figure_references=[])
                for n in re.findall(r"=== Slide (\d+) ===", section_text)],
        sections=[])
