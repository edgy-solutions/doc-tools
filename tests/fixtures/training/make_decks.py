"""Build 20 synthetic training decks (edge cases for the training-workbench contract).

    python make_decks.py <out_dir>   # run through the doc-tools repo venv only

Writes <out_dir>/<slug>/<file>.pptx and <out_dir>/<slug>/metadata.json (omitted
for empty-no-metadata). Content is deterministic; BYTES are not (zip member
timestamps), exactly as with tests/fixtures/training/make_deck.py, so a re-run
is a new set of ingest ids. Only dup-same-bytes-a/b are byte-identical to each
other (generated once, bytes copied). Images are drawn in memory with Pillow.
"""
import datetime
import io
import json
import sys
from pathlib import Path

from PIL import Image, ImageDraw
from pptx import Presentation
from pptx.util import Inches

FIXED = datetime.datetime(2026, 10, 23, 0, 0, 0)
W, H = Inches(13.333), Inches(7.5)
L_TITLE, L_BODY, L_SECTION, L_TITLE_ONLY, L_BLANK = 0, 1, 2, 5, 6


# ---------------------------------------------------------------- helpers
def png(label: str, seed: int = 0, size=(640, 360)) -> io.BytesIO:
    w, h = size
    img = Image.new("RGB", size)
    d = ImageDraw.Draw(img)
    base = ((seed * 53) % 200, (seed * 97 + 60) % 200, (seed * 31 + 120) % 200)
    for y in range(h):
        f = y / h
        d.line((0, y, w, y), fill=tuple(int(c + (255 - c) * f * 0.6) for c in base))
    d.rectangle((w // 8, h // 4, w // 2, h // 2), outline="black", width=4)
    d.ellipse((w // 2 + 20, h // 3, w - w // 8, h - h // 6), outline="black", width=4)
    d.line((w // 2, h * 3 // 8, w // 2 + 20, h // 2 + 20), fill="black", width=4)
    d.text((12, 12), label, fill="black")
    out = io.BytesIO()
    img.save(out, format="PNG")
    out.seek(0)
    return out


def new_prs(title: str) -> Presentation:
    prs = Presentation()
    prs.slide_width, prs.slide_height = W, H
    cp = prs.core_properties
    cp.title, cp.author = title, "doc-tools fixture"
    cp.created = cp.modified = cp.last_modified_by = FIXED
    cp.last_modified_by = "doc-tools fixture"
    return prs


def title_slide(prs, title, sub=""):
    s = prs.slides.add_slide(prs.slide_layouts[L_TITLE])
    s.shapes.title.text = title
    s.placeholders[1].text = sub
    return s


def section(prs, num, title):
    s = prs.slides.add_slide(prs.slide_layouts[L_SECTION])
    s.shapes.title.text = f"{num} {title}"
    return s


def content(prs, title, bullets):
    s = prs.slides.add_slide(prs.slide_layouts[L_BODY])
    s.shapes.title.text = title
    body = s.placeholders[1].text_frame
    body.text = bullets[0]
    for b in bullets[1:]:
        body.add_paragraph().text = b
    return s


def table_slide(prs, title, rows):
    s = prs.slides.add_slide(prs.slide_layouts[L_TITLE_ONLY])
    s.shapes.title.text = title
    gt = s.shapes.add_table(len(rows), len(rows[0]), Inches(1), Inches(2.2),
                            Inches(11), Inches(0.5) * len(rows))
    for r, row in enumerate(rows):
        for c, v in enumerate(row):
            gt.table.cell(r, c).text = str(v)
    return s


def blank(prs):
    return prs.slides.add_slide(prs.slide_layouts[L_BLANK])


def pic(s, label, seed, left, top, width):
    s.shapes.add_picture(png(label, seed), left, top, width=width)


def picture_slide(prs, label, seed=1):
    s = blank(prs)
    pic(s, label, seed, Inches(3.2), Inches(1.5), Inches(7))
    return s


def agenda(prs, outline):
    s = prs.slides.add_slide(prs.slide_layouts[L_BODY])
    s.shapes.title.text = "Agenda"
    tf = s.placeholders[1].text_frame
    tf.text = outline[0]
    for o in outline[1:]:
        tf.add_paragraph().text = o
    for p in tf.paragraphs:
        for r in p.runs:
            r.font.size = Inches(0.2)
    return s


def to_bytes(prs) -> bytes:
    b = io.BytesIO()
    prs.save(b)
    return b.getvalue()


def meta(title, bu, disc, aud, lvl, delivery, hours, ver):
    return {"title": title, "business_unit": bu, "engineering_discipline": disc,
            "audience": aud, "level_of_material": lvl,
            "current_delivery_method": delivery, "duration_hours": hours,
            "version": ver}


TOPICS = [
    ("Hydraulic Accumulators", ["Store energy as pressurised gas against fluid", "Bladder, piston and diaphragm types",
                                "Pre-charge is checked with the system depressurised", "Isolate and bleed before servicing"]),
    ("Bus Bar Protection", ["Differential schemes compare current in and out", "Zone boundaries follow the CT locations",
                            "Check CT polarity during commissioning"]),
    ("ARINC 429 Basics", ["Unidirectional serial bus, one transmitter", "32-bit words with label, SDI, data, SSM, parity",
                          "Low speed 12.5 kbit/s, high speed 100 kbit/s", "Up to 20 receivers per transmitter"]),
    ("Lockout and Tagout", ["Identify every energy source", "Apply personal locks and tags", "Verify zero energy before work",
                            "Only the owner removes a lock"]),
    ("Insulation Resistance Testing", ["Megger test at 500 V or 1000 V", "Record temperature with the reading",
                                       "Discharge the circuit after the test"]),
    ("Fuel Quantity Probes", ["Capacitance varies with fuel level", "Compensator corrects for density", "Wiring is intrinsically safe"]),
    ("Torque Wrench Use", ["Select a wrench in the middle of its range", "Never use as a breaker bar",
                           "Store at the lowest setting", "Calibrate on schedule"]),
    ("Three-Phase Power", ["Phases are 120 degrees apart", "Line voltage is root three times phase voltage",
                           "Balanced loads leave no neutral current"]),
    ("Confined Space Entry", ["Test the atmosphere before entry", "Attendant remains outside at all times",
                              "Rescue plan is in place before entry", "Permit expires at end of shift"]),
    ("Pitot-Static Systems", ["Pitot measures total pressure", "Static ports measure ambient pressure",
                              "Leak checks follow any disturbance of the lines"]),
    ("Relief Valve Setting", ["Set below the weakest component rating", "Verify with a calibrated gauge",
                              "Seal the adjustment after setting"]),
    ("Battery Maintenance", ["Check electrolyte level on vented cells", "Keep terminals clean and torqued",
                             "Never mix cell chemistries", "Ventilate during charging"]),
]


def topic(i):
    t, b = TOPICS[i % len(TOPICS)]
    return t, b


def add_topic_slides(prs, n, offset=0):
    for k in range(n):
        t, b = topic(offset + k)
        content(prs, t, b[: 3 + (offset + k) % 3] if len(b) >= 3 else b)


# ---------------------------------------------------------------- decks
def d01():
    p = new_prs("Nested Three Levels")
    title_slide(p, "Power Systems Overview", "Three-level section hierarchy")
    agenda(p, ["1 Generation", "1.1 Generators", "1.1.1 Synchronous Machines", "1.1.2 Excitation",
               "1.2 Prime Movers", "2 Distribution", "2.1 Switchgear", "2.1.1 Circuit Breakers"])
    section(p, "1", "Generation"); section(p, "1.1", "Generators"); section(p, "1.1.1", "Synchronous Machines")
    content(p, "Rotor and Stator", ["Rotor carries the field winding", "Stator carries the armature winding", "Speed sets frequency"])
    section(p, "1.1.2", "Excitation")
    content(p, "Excitation Systems", ["Static excitation uses a rectifier", "Brushless excitation uses a rotating diode bridge"])
    section(p, "1.2", "Prime Movers")
    content(p, "Turbine Types", ["Steam turbines", "Gas turbines", "Hydro turbines"])
    section(p, "2", "Distribution"); section(p, "2.1", "Switchgear"); section(p, "2.1.1", "Circuit Breakers")
    content(p, "Breaker Ratings", ["Rated voltage and current", "Short-circuit interrupting capacity"])
    return p


def d02():
    p = new_prs("Nested Four Levels")
    title_slide(p, "Aircraft Hydraulic Systems", "Four-level section hierarchy")
    agenda(p, ["1 Hydraulic Systems", "1.1 Power Generation", "1.1.1 Pumps", "1.1.1.1 Engine-Driven Pumps",
               "1.1.1.2 Electric Pumps", "1.1.2 Reservoirs", "1.2 Distribution", "1.2.1 Valves", "1.2.1.1 Selector Valves",
               "2 Maintenance", "2.1 Servicing", "2.1.1 Fluid Handling"])
    section(p, "1", "Hydraulic Systems"); section(p, "1.1", "Power Generation"); section(p, "1.1.1", "Pumps")
    section(p, "1.1.1.1", "Engine-Driven Pumps")
    content(p, "EDP Operation", ["Variable displacement piston pump", "Driven from the accessory gearbox", "Case drain returns to reservoir"])
    section(p, "1.1.1.2", "Electric Pumps")
    content(p, "Electric Motor Pumps", ["Backup when engines are shut down", "Duty-cycle limited"])
    section(p, "1.1.2", "Reservoirs")
    content(p, "Reservoir Pressurisation", ["Bleed-air pressurised", "Prevents pump cavitation at altitude"])
    section(p, "1.2", "Distribution"); section(p, "1.2.1", "Valves"); section(p, "1.2.1.1", "Selector Valves")
    content(p, "Selector Valve Types", ["Spool type", "Rotary type", "Solenoid operated"])
    section(p, "2", "Maintenance"); section(p, "2.1", "Servicing"); section(p, "2.1.1", "Fluid Handling")
    content(p, "Skydrol Precautions", ["Wear eye protection and nitrile gloves", "Keep away from paint and rubber seals"])
    return p


def d03():
    p = new_prs("Nested With Empty Sections")
    title_slide(p, "Electrical Safety Programme", "Contains sections with no content")
    agenda(p, ["1 Hazards", "1.1 Shock", "1.2 Arc Flash", "2 Controls", "2.1 PPE", "3 Appendix", "4 References"])
    section(p, "1", "Hazards")
    section(p, "1.1", "Shock")
    content(p, "Shock Effects", ["Current through the body is what harms", "Path matters as much as magnitude"])
    section(p, "1.2", "Arc Flash")  # empty: next slide is a section header
    section(p, "2", "Controls")
    section(p, "2.1", "PPE")
    content(p, "Arc-Rated Clothing", ["Rated in cal/cm2", "Layering adds protection"])
    section(p, "3", "Appendix")  # empty
    section(p, "4", "References")  # final, nothing after
    return p


def d04():
    p = new_prs("Nested Sibling Same Title")
    title_slide(p, "Avionics Fundamentals", "Repeated child titles under different parents")
    agenda(p, ["1 Communication Systems", "1.1 Summary", "1.2 Review Questions", "2 Navigation Systems",
               "2.1 Summary", "2.2 Review Questions"])
    section(p, "1", "Communication Systems")
    content(p, "VHF Communication", ["118 to 137 MHz AM voice", "Line-of-sight range"])
    section(p, "1.1", "Summary")
    content(p, "Communication Key Points", ["VHF for line of sight", "HF for oceanic"])
    section(p, "1.2", "Review Questions")
    content(p, "Communication Questions", ["What band does VHF use?", "Why is HF used over oceans?"])
    section(p, "2", "Navigation Systems")
    content(p, "VOR and DME", ["VOR gives bearing", "DME gives slant range"])
    section(p, "2.1", "Summary")
    content(p, "Navigation Key Points", ["VOR plus DME fixes position", "GNSS augments both"])
    section(p, "2.2", "Review Questions")
    content(p, "Navigation Questions", ["What does DME measure?", "What is slant range?"])
    return p


def d05():
    p = new_prs("Image Only Slide")
    title_slide(p, "Pump Teardown Walkthrough", "Slide 2 is a lone picture")
    picture_slide(p, "exploded view of gear pump", 5)
    content(p, "Inspection Points", ["Gear faces for scoring", "Bushings for wear", "Seals for hardening"])
    content(p, "Reassembly", ["Lubricate with system fluid", "Torque in a cross pattern"])
    content(p, "Close", ["Leak test at working pressure"])
    return p


def d06():
    p = new_prs("Image Multi Per Slide")
    title_slide(p, "Connector Identification Guide", "Contact sheets")
    s = blank(p)
    for i in range(4):
        pic(s, f"connector {i + 1}", 10 + i, Inches(0.5 + (i % 2) * 6.3), Inches(0.4 + (i // 2) * 3.4), Inches(6))
    content(p, "Notes", ["Match keying before mating"])
    s = blank(p)
    for i in range(6):
        pic(s, f"pin layout {i + 1}", 20 + i, Inches(0.4 + (i % 3) * 4.2), Inches(0.5 + (i // 3) * 3.4), Inches(4))
    return p


def d07():
    p = new_prs("Image Full Bleed")
    title_slide(p, "Hangar Safety Tour", "Full-bleed pictures")
    s = blank(p)
    s.shapes.add_picture(png("full bleed hangar", 7, (1280, 720)), 0, 0, width=W, height=H)
    s = prs_title_only = p.slides.add_slide(p.slide_layouts[L_TITLE_ONLY])
    s.shapes.add_picture(png("full bleed with title", 8, (1280, 720)), 0, 0, width=W, height=H)
    s.shapes.title.text = "Emergency Exits"
    # picture was added after the title placeholder, so it covers it; fine for the case
    content(p, "Assembly Points", ["North apron", "South apron"])
    return p


def d08():
    p = new_prs("Image Only Deck")
    for i in range(5):
        picture_slide(p, f"wiring diagram sheet {i + 1}", 30 + i)
    return p


def _long(title, n, gap=20, start_secs=True):
    p = new_prs(title)
    title_slide(p, title, f"{n} slides")
    k = 1
    while len(p.slides) < n:
        if (len(p.slides) - 1) % gap == 0:
            section(p, str(k), f"Module {k}")
            k += 1
        else:
            t, b = topic(len(p.slides))
            content(p, f"{t} ({len(p.slides)})", b[: 3 + len(p.slides) % 3])
    return p


def d09():
    return _long("Long Deck 160", 160)


def d10():
    p = _long("Long Deck With Blank Gaps", 155)
    blanks = {10, 50, 51, 120}
    rebuilt = new_prs("Long Deck With Blank Gaps")
    for i, src in enumerate(p.slides, 1):
        if i in blanks:
            blank(rebuilt)
        else:
            _clone(rebuilt, src)
    return rebuilt


def _clone(dst, src):
    """Re-create a title/body slide by layout name and text."""
    layout = {l.name: l for l in dst.slide_layouts}[src.slide_layout.name]
    s = dst.slides.add_slide(layout)
    for sh_src, sh_dst in zip(src.placeholders, s.placeholders):
        sh_dst.text_frame.text = sh_src.text_frame.text
        if sh_src.text_frame.text and len(sh_src.text_frame.paragraphs) > 1:
            tf = sh_dst.text_frame
            tf.text = sh_src.text_frame.paragraphs[0].text
            for para in sh_src.text_frame.paragraphs[1:]:
                tf.add_paragraph().text = para.text


def d11():
    p = _long("Long Deck With Hidden Slides", 150)
    for n in (30, 31, 90):
        p.slides[n - 1]._element.set("show", "0")
    return p


def d12():
    p = new_prs("Pump Maintenance Basics")
    title_slide(p, "Pump Maintenance Basics", "Edition A: mechanical seals")
    content(p, "Mechanical Seals", ["Faces lapped flat", "Never run dry", "Replace as a set"])
    content(p, "Seal Failure Modes", ["Dry running", "Contamination", "Thermal shock"])
    return p


def d13():
    p = new_prs("Pump Maintenance Basics")
    title_slide(p, "Pump Maintenance Basics", "Edition B: bearings and alignment")
    content(p, "Bearing Lubrication", ["Grease quantity matters", "Over-greasing raises temperature"])
    content(p, "Shaft Alignment", ["Dial indicator method", "Laser alignment", "Soft foot check first"])
    table_slide(p, "Alignment Tolerances", [["Speed (rpm)", "Offset (mm)"], ["1800", "0.05"], ["3600", "0.03"]])
    return p


def d14():
    p = new_prs("Concept Spelling Variants")
    title_slide(p, "Hydraulic Pump Fundamentals", "Same concepts, spelled several ways")
    content(p, "Hydraulic Pump Types", ["Hydraulic Pump: gear, vane, piston", "hydraulic pump selection depends on pressure"])
    content(p, "HYDRAULIC PUMP SAFETY", ["HYDRAULIC PUMP casings can be hot", "Hydraulic-Pump inlets must not be starved"])
    content(p, "Pressure Relief Valve Role", ["Pressure Relief Valve protects the circuit",
                                              "pressure-relief valve cracking pressure is set below the maximum"])
    content(p, "Relief Valve Testing", ["Pressure relief valve(s) are bench tested", "A Hydraulic Pump. feeds the relief valve under test"])
    content(p, "Summary", ["Hydraulic Pump", "Pressure Relief Valve", "hydraulic pump and pressure-relief valve work as a pair"])
    return p


def d15():
    p = new_prs("Identical Bytes Course")
    title_slide(p, "Corrosion Control", "Byte-identical copy test")
    content(p, "Corrosion Types", ["Galvanic", "Pitting", "Crevice"])
    content(p, "Prevention", ["Primers and sealants", "Drain holes kept clear"])
    return p


def d17():
    p = new_prs("Empty Titles")
    title_slide(p, "Gearbox Inspection", "Some slides have empty titles")
    s = p.slides.add_slide(p.slide_layouts[L_BODY])  # title placeholder present, empty
    s.placeholders[1].text_frame.text = "Check the magnetic chip detector"
    s = p.slides.add_slide(p.slide_layouts[L_TITLE_ONLY])  # empty title, nothing else
    s = p.slides.add_slide(p.slide_layouts[L_BODY])
    s.shapes.title.text = "   "  # whitespace-only title
    s.placeholders[1].text_frame.text = "Oil sample analysis"
    content(p, "Wrap-up", ["Record all findings"])
    return p


def d18():
    p = new_prs("Empty Blank Slide")
    title_slide(p, "Torque Tools", "A blank slide mid-deck and last")
    content(p, "Click Wrenches", ["Set to value", "Return to minimum"])
    blank(p)
    content(p, "Electronic Wrenches", ["Log each fastener", "Angle control"])
    blank(p)
    return p


def d20():
    p = new_prs("Whitespace And Notes")
    title_slide(p, "Instrument Calibration", "Whitespace-only text and notes-only text")
    s = p.slides.add_slide(p.slide_layouts[L_TITLE_ONLY])
    s.shapes.title.text = "Reference Standards"
    for i, txt in enumerate(["   ", "\n\n", "\t \t", " \n \t "]):
        tb = s.shapes.add_textbox(Inches(1), Inches(2 + i * 0.7), Inches(6), Inches(0.5))
        tb.text_frame.text = txt
    s = blank(p)
    s.notes_slide.notes_text_frame.text = "Speaker note only: explain the calibration chain from the working standard to the national standard."
    s = content(p, "Calibration Intervals", ["Annual by default", "Shorter after a drop or overload"])
    s.notes_slide.notes_text_frame.text = "Mention the drift trend chart."
    return p


# ---------------------------------------------------------------- main
M = meta
DECKS = [
    ("nested-three-levels", d01, M("Power Systems Overview", "Energy Systems", "Electrical", "Plant operators", "Intermediate", "Instructor-led", 3, "1.0")),
    ("nested-four-levels", d02, M("Aircraft Hydraulic Systems", "Aerospace Services", "Mechanical", "Line mechanics", "Intermediate", "Blended", 4, "2.1")),
    ("nested-empty-section", d03, M("Electrical Safety Programme", "Facilities", "Electrical", "All technicians", "Introductory", "Instructor-led", 1.5, "1.2")),
    ("nested-sibling-same-title", d04, M("Avionics Fundamentals", "Aerospace Services", "Avionics", "Avionics technicians", "Introductory", "E-learning", 2, "3.0")),
    ("image-only-slide", d05, M("Pump Teardown Walkthrough", "Industrial Pumps", "Mechanical", "Workshop fitters", "Intermediate", "Instructor-led", 1, "1.0")),
    ("image-multi-per-slide", d06, M("Connector Identification Guide", "Avionics Supply", "Electrical", "Stores staff", "Introductory", "Self-study", 0.5, "1.1")),
    ("image-full-bleed", d07, M("Hangar Safety Tour", "Facilities", "Safety", "New starters", "Introductory", "Instructor-led", 0.75, "1.0")),
    ("image-only-deck", d08, M("Wiring Diagram Pack", "Avionics Supply", "Electrical", "Troubleshooters", "Advanced", "Self-study", 1, "4.2")),
    ("long-160-slides", d09, M("Maintenance Technician Core Course", "Maintenance Training", "Mechanical", "Maintenance technicians", "Intermediate", "Instructor-led", 40, "5.0")),
    ("long-with-blank-gaps", d10, M("Systems Familiarisation Course", "Maintenance Training", "Systems", "Flight-line engineers", "Intermediate", "Blended", 38, "2.3")),
    ("long-with-hidden-slides", d11, M("Refresher Course With Hidden Slides", "Maintenance Training", "Systems", "Certifying staff", "Advanced", "Blended", 35, "1.4")),
    ("dup-title-a", d12, M("Pump Maintenance Basics", "Industrial Pumps", "Mechanical", "Fitters", "Introductory", "Instructor-led", 2, "1.0")),
    ("dup-title-b", d13, M("Pump Maintenance Basics", "Industrial Pumps", "Mechanical", "Apprentices", "Introductory", "E-learning", 2.5, "2.0")),
    ("dup-concept-case", d14, M("Hydraulic Pump Fundamentals", "Fluid Power", "Hydraulics", "Maintenance technicians", "Introductory", "Instructor-led", 1, "1.0")),
    ("dup-same-bytes-a", d15, M("Corrosion Control", "Structures Repair", "Materials", "Structures technicians", "Intermediate", "Instructor-led", 2, "1.0")),
    ("dup-same-bytes-b", None, None),
    ("empty-title", d17, M("Gearbox Inspection", "Rotorcraft Services", "Mechanical", "Inspectors", "Intermediate", "Instructor-led", 1, "1.0")),
    ("empty-blank-slide", d18, M("Torque Tools", "Tooling", "Mechanical", "Workshop staff", "Introductory", "Self-study", 0.5, "1.0")),
    ("empty-no-metadata", d17, None),
    ("empty-whitespace-and-notes", d20, M("Instrument Calibration", "Metrology", "Instrumentation", "Calibration technicians", "Advanced", "Instructor-led", 3, "1.3")),
]
FILES = {"dup-same-bytes-a": "same_bytes_a.pptx", "dup-same-bytes-b": "same_bytes_b.pptx"}


def main(out: Path) -> None:
    same_bytes = None
    same_meta = None
    for slug, fn, md in DECKS:
        d = out / slug
        d.mkdir(parents=True, exist_ok=True)
        name = FILES.get(slug, f"{slug}.pptx")
        if slug == "dup-same-bytes-a":
            same_bytes, same_meta = to_bytes(fn()), md
            data, md = same_bytes, same_meta
        elif slug == "dup-same-bytes-b":
            data, md = same_bytes, same_meta
        elif slug == "empty-no-metadata":
            data = to_bytes(d19())
        else:
            data = to_bytes(fn())
        (d / name).write_bytes(data)
        if md is not None:
            (d / "metadata.json").write_text(json.dumps(md, indent=2) + "\n", encoding="utf-8")
        print(slug, name, len(data))


def d19():
    p = new_prs("Deck Without Sidecar")
    title_slide(p, "Rigging Basics", "No metadata.json accompanies this deck")
    content(p, "Sling Angles", ["Lower angles raise sling tension", "Never exceed the rated angle"])
    content(p, "Inspection", ["Check tags and stitching", "Reject any damaged sling"])
    return p


if __name__ == "__main__":
    main(Path(sys.argv[1]))
