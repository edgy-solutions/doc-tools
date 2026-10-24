"""Regenerate the training-workbench fixture deck (contract fixture).

The committed ``hydraulics_fixture.pptx`` is canonical; its sha256 is pinned in
``golden.json`` and in the workbench's vendored copy. This script records how it
was made. Re-running it yields the same CONTENT but not the same BYTES (zip
member timestamps), so a re-run is a contract change: re-pin the sha256 and send
wb a packet. Run through the repo venv only:

    .venv/Scripts/python tests/fixtures/training/make_deck.py
"""
from pathlib import Path

from PIL import Image, ImageDraw
from pptx import Presentation
from pptx.util import Inches

HERE = Path(__file__).parent


def _schematic(path: Path) -> None:
    img = Image.new("RGB", (480, 240), "white")
    d = ImageDraw.Draw(img)
    d.rectangle((20, 80, 140, 200), outline="black", width=4)    # reservoir
    d.rectangle((190, 110, 290, 170), outline="black", width=4)  # filter
    d.ellipse((340, 90, 440, 190), outline="black", width=4)     # pump
    d.line((140, 140, 190, 140), fill="black", width=4)
    d.line((290, 140, 340, 140), fill="black", width=4)
    img.save(path, format="PNG")


def main() -> None:
    prs = Presentation()
    title = prs.slides.add_slide(prs.slide_layouts[0])
    title.shapes.title.text = "Hydraulic Systems Fundamentals"
    title.placeholders[1].text = "Module 1: fixture deck for the training-workbench contract"

    for heading, bullets in [
        ("Pascal's Law", [
            "Pressure applied to a confined fluid is transmitted equally in all directions",
            "Force equals pressure multiplied by piston area",
            "A small input piston can lift a large load",
        ]),
        ("Hydraulic Pumps", [
            "Gear pumps are simple and tolerant of contamination",
            "Vane pumps run quietly at moderate pressure",
            "Piston pumps deliver the highest pressure",
        ]),
    ]:
        s = prs.slides.add_slide(prs.slide_layouts[1])
        s.shapes.title.text = heading
        body = s.placeholders[1].text_frame
        body.text = bullets[0]
        for b in bullets[1:]:
            body.add_paragraph().text = b

    s = prs.slides.add_slide(prs.slide_layouts[5])  # title only
    s.shapes.title.text = "Reservoir Schematic"
    png = HERE / "_schematic.png"
    _schematic(png)
    s.shapes.add_picture(str(png), Inches(1.5), Inches(2.0), width=Inches(6))
    tb = s.shapes.add_textbox(Inches(1.5), Inches(5.5), Inches(6), Inches(0.6))
    tb.text_frame.text = "Figure 1: reservoir, filter and pump in series"
    png.unlink()

    prs.save(HERE / "hydraulics_fixture.pptx")


if __name__ == "__main__":
    main()
